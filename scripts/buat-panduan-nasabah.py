"""Generate docs/paparan-nasabah.pptx dari naskah slide di bawah (mirror docs/nasabah.md)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent.parent / "docs" / "paparan-nasabah.pptx"
ACCENT = "1D4ED8"  # biru brand

SLIDES = [
    ("Panduan Nasabah — Bank Miska", [
        "Untuk siswa pemilik rekening: login, daftar, 3 menu, password, bantuan.",
        "Syarat: sudah punya rekening (kode Hxxx tabungan / Pxxx piutang).",
        "Demo: siswa / siswa123.",
    ]),
    ("1. Masuk (login)", [
        "Buka halaman Masuk, isi username + password, klik Masuk.",
        "Sesi berlaku 7 hari; bila habis, login ulang.",
        "Gagal? Cek pesan: salah kredensial, akun dinonaktifkan, atau DB tidak terhubung.",
    ]),
    ("2. Daftar akun baru", [
        "Khusus yang sudah punya rekening: halaman Daftar.",
        "Isi username (3-50 karakter), password (min 6), kode rekening Hxxx/Pxxx.",
        "Akun BELUM AKTIF sampai diverifikasi teller/admin di loket.",
        "Satu kode rekening untuk satu akun.",
    ]),
    ("3. Menu: Riwayat Saya", [
        "Mutasi + saldo rekening SENDIRI (server memaksa kode milik sendiri).",
        "Filter rentang tanggal, lihat saldo akhir + total masuk/keluar.",
        "Unduh CSV / XLSX.",
        '"Akun belum ditautkan" = minta teller menautkan kode.',
    ]),
    ("4. Menu: Daftar Saldo", [
        "Saldo per tipe rekening: TABUNGAN atau PIUTANG.",
        "Tabel kode + nama + saldo akhir + total, bisa diekspor.",
    ]),
    ("5. Password & keluar", [
        "Sudah login: ganti sendiri di Ubah Password (perlu password lama).",
        "Belum bisa login: datang ke loket, admin meresetkan (ada popup konfirmasi).",
        "Keluar: klik nama di bawah sidebar, lalu Keluar.",
    ]),
    ("6. Batasan & bantuan", [
        "Nasabah hanya baca data sendiri; halaman staf = Akses ditolak (normal).",
        "Akun nonaktif = hubungi admin.",
        "Panduan lengkap: docs/nasabah.md; bantuan dalam aplikasi: menu Bantuan.",
    ]),
]


def main():
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    for judul, bullets in SLIDES:
        s = prs.slides.add_slide(blank)
        tx = s.shapes.add_textbox(Emu(600000), Emu(300000), Emu(10992000), Emu(900000)).text_frame
        tx.word_wrap = True
        p = tx.paragraphs[0]
        p.text = judul
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor.from_string(ACCENT)
        bx = s.shapes.add_textbox(Emu(600000), Emu(1500000), Emu(10992000), Emu(4500000)).text_frame
        bx.word_wrap = True
        for i, b in enumerate(bullets):
            p = bx.paragraphs[0] if i == 0 else bx.add_paragraph()
            p.text = "•  " + b
            p.font.size = Pt(20)
            p.space_after = Pt(10)
        ft = s.shapes.add_textbox(Emu(600000), Emu(6200000), Emu(10992000), Emu(400000)).text_frame
        ft.word_wrap = True
        fp = ft.paragraphs[0]
        fp.text = "Bank Miska — panduan nasabah"
        fp.font.size = Pt(12)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print("OK:", OUT, OUT.stat().st_size, "bytes,", len(SLIDES), "slide")


if __name__ == "__main__":
    main()
