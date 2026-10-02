"""Generate docs/paparan-investor.pptx dari naskah slide di bawah (mirror paparan-investor.md)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent.parent / "docs" / "paparan-investor.pptx"
ACCENT = "1D4ED8"  # biru brand

SLIDES = [
    ("Bank Miska", [
        "Kasir cashless + mini-bank sekolah: tap kartu / scan QR, jurnal otomatis, laporan selalu balance.",
        "Status: berjalan — 18 halaman, ~30 API, ekspor Excel 17-sheet gaya workbook sekolah.",
        "Dokumen rahasia investor.",
    ]),
    ("1. Masalah", [
        "Kantin tunai: antre istirahat panjang, kembalian lambat, selisih kas sulit ditelusur.",
        "Excel rapuh: #N/A VLOOKUP putus, CIRCULAR REFERENCES, SUM manual, saldo awal fiktif (kasus nyata Kas −Rp 99 jt).",
        "Laporan gaya workbook sekolah makan waktu berminggu-minggu tiap tutup bulan.",
    ]),
    ("2. Solusi (demo 10 detik)", [
        "Top-up sekali ke teller → jajan cukup tap RFID / scan QR 5-menit.",
        "Setiap transaksi = jurnal berpasangan otomatis (Kas, Tabungan, Piutang, Bunga).",
        "Salah input? Buka periode → jurnal koreksi → tutup ulang; arsip bulan bisa dicetak ulang kapan pun.",
    ]),
    ("3. Konsep desain (lihat XMind)", [
        "Pengguna: teller siswa, admin guru, nasabah Hxxx/Pxxx.",
        "Transaksi → Akuntansi (jurnal → BB → NL → penutup → arsip).",
        "Teknologi / Keamanan / Operasional / Roadmap (detail di mindmap).",
    ]),
    ("4. Teknologi", [
        "UI: Next.js 16 + React 19 + Tailwind — 18 halaman loket & HP.",
        "Backend: Route Handlers TypeScript + mysql2 tanpa ORM — tiap aksi uang = transaksi DB atomik.",
        "Database: MySQL skema tunggal; RFID wedge + QR, exceljs untuk laporan.",
    ]),
    ("5. Keamanan 5 lapis", [
        "L1 bcrypt + sesi HttpOnly 7 hari; L2 otorisasi role per API.",
        "L3 anti-overdraft, anti-Kas-negatif, cair wajib ke Pxxx.",
        "L4 kunci periode + buka-terbaru; L5 tiap rupiah tertelusur ke jurnal.",
        "Jujur: rate-limit, audit log admin, PIN nasabah = roadmap.",
    ]),
    ("6. Target & kapasitas (ESTIMASI, 10.000 pengguna)", [
        "40% jajan/hari × 2 trx = ±8.000 transaksi/hari.",
        "Puncak istirahat ±1.500 trx/jam (≈ 1 trx / 2,4 dtk).",
        "Desain: pool MySQL 10, tulis atomik < 200 ms — ampun untuk puncak.",
        "Validasi 2 minggu pilot sebelum kunci angka.",
    ]),
    ("7. Panel admin operator", [
        "Pengguna: tambah/nonaktif, pasang-lepas UID kartu.",
        "Merchant: kelola kantin + QR, pantau status bayar.",
        "Periode + arsip: tutup/buka, ekspor paket Excel per bulan.",
        "Rencana: limit harian, audit log, dashboard puncak.",
    ]),
    ("8. Maintenance tahunan (ESTIMASI/sekolah)", [
        "VPS + backup Rp 3–6 jt; domain+SSL Rp 0,3–0,5 jt; RFID Rp 2–5 jt.",
        "Support SLA Rp 6–12 jt; retainer dev Rp 12–24 jt.",
        "Total ±Rp 24–48 jt (< Rp 5 rb/pengguna/tahun).",
    ]),
    ("9. Model bisnis (opsi)", [
        "SaaS per sekolah menutup maintenance + margin.",
        "Setup awal: instalasi + reader + training (satu kali).",
        "Skala: kode sama, multitenant tahap 2.",
    ]),
    ("10. Roadmap", [
        "Q1 transfer antar-nasabah satu sekolah + reversal kantin.",
        "Q2 hardening: rate-limit, audit log, PIN.",
        "Q3 snapshot arsip + dashboard; Q4 multitenant.",
    ]),
    ("11. Ask & penutup", [
        "Butuh: pilot 1 SMK (10.000 pengguna) + dana setup & hardening Q1–Q2.",
        "Tawaran: demo live, arsip Excel contoh, mindmap konsep.",
        "Kontak: [isi nama / no / WA].",
    ]),
]


def main():
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    for judul, bullets in SLIDES:
        s = prs.slides.add_slide(blank)
        tx = s.shapes.add_textbox(Emu(600000), Emu(300000), Emu(10992000), Emu(900000)).text_frame
        tx.word_wrap = True
        p = tx.paragraphs[0]
        p.text = judul
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor.from_string(ACCENT)
        bx = s.shapes.add_textbox(Emu(600000), Emu(1500000), Emu(10992000), Emu(4500000)).text_frame
        bx.word_wrap = True
        for i, b in enumerate(bullets):
            p = bx.paragraphs[0] if i == 0 else bx.add_paragraph()
            p.text = "•  " + b
            p.font.size = Pt(20)
            p.space_after = Pt(10)
        ft = s.shapes.add_textbox(Emu(600000), Emu(6200000), Emu(10992000), Emu(400000)).text_frame
        ft.word_wrap = True
        fp = ft.paragraphs[0]
        fp.text = "Bank Miska — rahasia investor"
        fp.font.size = Pt(12)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print("OK:", OUT, OUT.stat().st_size, "bytes,", len(SLIDES), "slide")


if __name__ == "__main__":
    main()
